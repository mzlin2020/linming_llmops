"""文件解析：把一个落盘文件按扩展名读成 LangChain Document 列表。

用轻量的 loader 组合（避免 unstructured 的重型依赖）：
- txt / md / markdown → TextLoader（md 直接按纯文本读，足够喂给切分器）
- csv                 → CSVLoader
- pdf                 → PyPDFLoader（依赖 pypdf）
- docx               → Docx2txtLoader（依赖 docx2txt）
- xlsx               → 自实现 openpyxl 逐单元格读取（openpyxl 只读 xlsx，不支持旧版二进制 xls）

扩展名→loader 的映射集中在 _LOADERS 一处，支持的扩展名 SUPPORTED_EXTENSIONS 由它派生（单一真相）。
所有 loader 延迟 import：缺某个解析包不影响其它格式，也不拖累模块导入。
返回 list[Document]；解析失败抛异常，由 IndexingService 捕获写入 document.error。
"""
from __future__ import annotations

import os
from typing import Callable, Dict, List

from langchain_core.documents import Document as LCDocument


def _normalize_ext(extension: str) -> str:
    return (extension or "").lower().lstrip(".")


def is_supported(extension: str) -> bool:
    return _normalize_ext(extension) in SUPPORTED_EXTENSIONS


def load(file_path: str, extension: str) -> List[LCDocument]:
    """读取 file_path（绝对路径）→ list[Document]。extension 为小写扩展名(不带点)。"""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")
    loader = _LOADERS.get(_normalize_ext(extension))
    if loader is None:
        raise ValueError(f"不支持的文件类型: {extension}")
    return loader(file_path)


def load_text(file_path: str, extension: str) -> str:
    """便捷：把解析结果拼成单个字符串（各 Document 的 page_content 用换行连接）。"""
    docs = load(file_path, extension)
    return "\n\n".join(d.page_content for d in docs if (d.page_content or "").strip())


# ---------------- 各格式 loader（延迟 import）----------------

def _load_text(file_path: str) -> List[LCDocument]:
    from langchain_community.document_loaders import TextLoader
    # autodetect_encoding：兼容非 utf-8 文本，避免 UnicodeDecodeError
    return TextLoader(file_path, autodetect_encoding=True).load()


def _load_csv(file_path: str) -> List[LCDocument]:
    from langchain_community.document_loaders import CSVLoader
    return CSVLoader(file_path, autodetect_encoding=True).load()


def _load_pdf(file_path: str) -> List[LCDocument]:
    from langchain_community.document_loaders import PyPDFLoader
    return PyPDFLoader(file_path).load()


def _load_docx(file_path: str) -> List[LCDocument]:
    from langchain_community.document_loaders import Docx2txtLoader
    return Docx2txtLoader(file_path).load()


def _load_excel(file_path: str) -> List[LCDocument]:
    """逐 sheet 逐行读非空单元格，拼成一份文本。"""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        lines: List[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(" ".join(cells))
    finally:
        wb.close()
    return [LCDocument(page_content="\n".join(lines), metadata={"source": file_path})]


# 扩展名 → loader 的单一真相；is_supported / SUPPORTED_EXTENSIONS 全部由它派生
_LOADERS: Dict[str, Callable[[str], List[LCDocument]]] = {
    "txt": _load_text,
    "md": _load_text,
    "markdown": _load_text,
    "csv": _load_csv,
    "pdf": _load_pdf,
    "docx": _load_docx,
    "xlsx": _load_excel,
}

# 支持的扩展名（小写，不带点）
SUPPORTED_EXTENSIONS = frozenset(_LOADERS)
